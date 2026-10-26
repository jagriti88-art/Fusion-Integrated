from docx import Document
from openpyxl import load_workbook
from pathlib import Path

root = Path.home() / "Downloads/ELM-Assignment-6/ELM-Assignment-6/ELM_Specifications"
out = Path("ELM_REQUIREMENTS.txt")

with out.open("w", encoding="utf-8") as f:
    for p in sorted(root.iterdir()):
        if p.suffix == ".docx":
            f.write(f"\n\n===== {p.name} =====\n")
            doc = Document(p)
            for para in doc.paragraphs:
                if para.text.strip():
                    f.write(para.text + "\n")
            for table in doc.tables:
                for row in table.rows:
                    f.write(" | ".join(c.text.replace("\n", " ") for c in row.cells) + "\n")
        elif p.suffix == ".xlsx":
            f.write(f"\n\n===== {p.name} =====\n")
            wb = load_workbook(p, read_only=True, data_only=True)
            for ws in wb.worksheets:
                f.write(f"\n--- Sheet: {ws.title} ---\n")
                for row in ws.iter_rows(values_only=True):
                    values = [str(v) if v is not None else "" for v in row]
                    if any(values):
                        f.write(" | ".join(values) + "\n")

print(f"Requirements extracted to {out.resolve()}")
